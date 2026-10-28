"""Fábrica de planilhas .xlsx sintéticas no formato da aba Base."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from opex_etl.domain import ABA_BASE, COLUNAS_BASE

LINHA_PADRAO: dict[str, object] = {
    "Nome da Origem": "GL_Realizado",
    "Cenário": "REALIZADO",
    "CEN2": "REAL",
    "Centro de Custo": "CC1",
    "Conta": "C1",
    "Projeto": "P1 - PROJETO UM/NA/NA",
    "Finalidade": "F1",
    "Grupo": "CONTRATOS",
    "Fornecedor": "FORN A",
    "Mês": "Janeiro",
    "Valor": 100.0,
    "DS_LINHA_LANCTO_CONTABIL": "LANC 1",
    "OC/RC": "SEM OC",
    "DT_CONTABILIZACAO": datetime(2026, 1, 31),
    "N4": "N4 A",
    "N5": "N5 A",
}

_NOME_ABA_ANALISE = "Análise"
_LINHA_CABECALHO_EXCEL = 2
_PRIMEIRA_LINHA_DADOS_EXCEL = 3


def criar_planilha_base(
    caminho: Path,
    linhas: list[dict[str, object]],
    colunas: list[str] | None = None,
    nome_aba: str = ABA_BASE,
) -> Path:
    colunas_planilha = list(COLUNAS_BASE) if colunas is None else colunas

    wb = Workbook()
    wb.active.title = _NOME_ABA_ANALISE  # type: ignore[union-attr]
    aba = wb.create_sheet(nome_aba)

    for indice, nome in enumerate(colunas_planilha, start=1):
        aba.cell(row=_LINHA_CABECALHO_EXCEL, column=indice, value=nome)

    for n_linha, dados in enumerate(linhas, start=_PRIMEIRA_LINHA_DADOS_EXCEL):
        for indice, nome in enumerate(colunas_planilha, start=1):
            valor = dados.get(nome)
            if valor is not None:
                aba.cell(row=n_linha, column=indice, value=valor)

    wb.save(caminho)
    wb.close()
    return caminho
